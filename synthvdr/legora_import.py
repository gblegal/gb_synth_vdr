"""Turn a Legora review into a tool output `python3 -m synthvdr score` can mark.

Legora reviews a blind room with the `review-a-test-data-room` skill and saves
a report in the shape that skill's `2-what-to-hand-back.md` fixes: one issue
per `##` heading, a severity line first, every document cited by its path, and
a `# Files read` list last. That is close to what `parse_markdown_report`
reads, but not close enough to score as it stands:

- Legora cites the path it saw. That starts at a mount such as
  `/workspace/documents/projects/<project name>/`, often names the tree the
  room was uploaded as, and on a rendered cut ends `.pdf` or `.docx`.
  `prematch` joins on exact strings against `.md` evidence paths, so a run
  over `data-room-pdf/` scores zero until every citation is cut back to the
  room-relative `.md` form. The cut is made at the first folder that is one
  of the room's SECTION_DIRS, which copes with any mount, any tree name and a
  project name with spaces in it — a space the parser's own path pattern
  cannot read.
- The `# Files read` list must never reach the parser. A level-1 heading does
  not start a finding, so the list would be read as the tail of the last
  issue, which would then cite every document in the room and pre-match every
  planted finding.
- Word drops the backticks, so a bare path counts as a citation when it
  starts at a section folder. A bare file name in prose ("README.md") does
  not; a backticked one does, so that a path cited without its folders
  reaches the room check and is refused rather than silently lost.

All of that leniency lives here; `score` stays strict and reads the result
as it reads any tool's JSON.
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Dict, List, Optional, Sequence, Set, Tuple

from .manifest import MANIFEST_NAME, read_content_hash
from .roomconf import RoomConf, load_room_conf
from .schema import SEVERITIES
from .score import ToolOutputError, parse_markdown_report


class LegoraImportError(Exception):
    """The report cannot be imported against this room — a path the uploaded
    cut does not hold, a report with no issues, an unreadable file, or an
    `--out` that would write into the room. Always a refusal, never a
    partial import: a scorecard built on it would not be a measurement of
    this room."""


# What a citation can name. `.pdf` and `.docx` are renders of a room's `.md`
# documents (`data-room-pdf/`, `data-room-docx/`), so they map back to it;
# `.csv` is a document in its own right, as the parser's own pattern allows.
DOCUMENT_SUFFIXES = (".md", ".pdf", ".docx", ".csv")
RENDERED_SUFFIXES = (".pdf", ".docx")

# How many paths a summary or a refusal names before it says "and N more".
LIST_LIMIT = 20

# The most documents one issue may cite before the import refuses it. The
# longest evidence chain in any room built so far is five documents; an issue
# citing dozens is almost always a list of paths read as part of it, which
# pre-matches every finding it touches. --allow-wide-issues overrides.
MAX_CITATIONS = 20


def as_markdown(rel: str) -> str:
    """A rendered document's path mapped back to the `.md` it was rendered from."""
    stem, dot, suffix = rel.rpartition(".")
    if dot and f".{suffix.lower()}" in RENDERED_SUFFIXES:
        return f"{stem}.md"
    return rel


def normalise_path(raw: str, section_dirs: Sequence[str]) -> str:
    """The room-relative `.md` path a cited path names, or the path as written
    when no component of it is a section folder — which then fails the room
    check, as it should."""
    sections = set(section_dirs)
    parts = [p for p in raw.strip().replace("\\", "/").split("/") if p not in ("", ".")]
    for index, part in enumerate(parts):
        if part in sections:
            return as_markdown("/".join(parts[index:]))
    return raw.strip()


# A bare path token, optionally wrapped in Markdown emphasis (`_x.pdf_`). Its
# edges exclude a trailing full stop so a path ending a sentence is still read,
# but not a dot followed by more name (".md.bak").
_BARE = r"(?<![\w./-])[*_]*([\w./-]+\.(?:md|pdf|docx|csv))[*_]*(?![\w/-]|\.\w)"
_BARE_PATH = re.compile(_BARE, re.IGNORECASE)
# A backticked span, or a bare path token.
_CITATION = re.compile(r"`([^`\n]+)`|" + _BARE, re.IGNORECASE)


def _unescape(text: str) -> str:
    """Markdown escapes a model or a Word export may put in a path: `01\\_corporate`."""
    return text.replace("\\_", "_").replace("\\*", "*")


def _anchored(raw: str, section_dirs: Sequence[str]) -> Optional[str]:
    """A bare token's normalised path, or None when it does not start at a
    section folder: a file name in prose, not a citation."""
    path = normalise_path(raw, section_dirs)
    return path if path.split("/", 1)[0] in set(section_dirs) else None


def _cited(match: "re.Match[str]", section_dirs: Sequence[str]) -> List[Tuple[str, str]]:
    """What one match cites, as (as written, normalised) pairs — none, one, or
    several.

    A backticked span that is a path is one citation, spaces and mount
    included. A span holding a path and more — `x.pdf#page=3`, `x.pdf, clause
    4.2`, `x.md:12` — is searched for bare paths, or the citation inside it is
    lost and its finding silently missed. A backticked document name that
    starts at no section folder is still returned, so the room check refuses
    it rather than the import dropping it unseen."""
    backticked, bare = match.group(1), match.group(2)
    if bare is not None:
        path = _anchored(bare, section_dirs)
        return [(bare, path)] if path else []
    path = normalise_path(backticked, section_dirs)
    if path.split("/", 1)[0] in set(section_dirs) and path.lower().endswith(DOCUMENT_SUFFIXES):
        return [(backticked, path)]
    inner = []
    for found in _BARE_PATH.finditer(backticked):
        anchored = _anchored(found.group(1), section_dirs)
        if anchored:
            inner.append((found.group(1), anchored))
    if inner:
        return inner
    if backticked.strip().lower().endswith(DOCUMENT_SUFFIXES):
        return [(backticked, path)]
    return []


def normalise_citations(text: str, section_dirs: Sequence[str]) -> str:
    """`text` with every citation rewritten as its normalised path in backticks,
    the form `parse_markdown_report` reads. Everything else is left as it is,
    bar the Markdown escapes inside paths."""

    def rewrite_inside(found: "re.Match[str]") -> str:
        anchored = _anchored(found.group(1), section_dirs)
        return found.group(0) if anchored is None else f"`{anchored}`"

    def replace(match: "re.Match[str]") -> str:
        cited = _cited(match, section_dirs)
        if not cited:
            return match.group(0)
        if match.group(1) is not None and len(cited) == 1 and cited[0][0] == match.group(1):
            return f"`{cited[0][1]}`"
        if match.group(1) is not None:
            # A span with more than a path: keep its words, lift the paths out.
            return _BARE_PATH.sub(rewrite_inside, match.group(1))
        return f"`{cited[0][1]}`"

    return _CITATION.sub(replace, _unescape(text))


def citation_pairs(text: str, section_dirs: Sequence[str]) -> List[Tuple[str, str]]:
    """Every citation in `text` as (as written, normalised), in order."""
    pairs: List[Tuple[str, str]] = []
    for match in _CITATION.finditer(_unescape(text)):
        pairs.extend(_cited(match, section_dirs))
    return pairs


def citations_in(text: str, section_dirs: Sequence[str]) -> List[str]:
    """Every document `text` cites, normalised, in the order cited."""
    return [path for _, path in citation_pairs(text, section_dirs)]


def room_trees(room: Path, section_dirs: Sequence[str]) -> Set[str]:
    """The room's top-level folders that hold section folders: every cut of it
    — data-room/, subset/, data-room-pdf/, data-room-docx/, a corrupted twin."""
    return {
        entry.name
        for entry in room.iterdir()
        if entry.is_dir()
        and not entry.name.startswith(".")
        and any((entry / section).is_dir() for section in section_dirs)
    }


def tree_named(raw: str, section_dirs: Sequence[str], trees: Set[str]) -> Optional[str]:
    """The room tree a cited path names just before its section folder, if any:
    `…/data-room-pdf/01_corporate/…` names data-room-pdf."""
    sections = set(section_dirs)
    by_fold = {tree.casefold(): tree for tree in trees}
    parts = [p for p in raw.strip().replace("\\", "/").split("/") if p]
    for index, part in enumerate(parts):
        if part in sections:
            return by_fold.get(parts[index - 1].casefold()) if index else None
    return None


def _within(path: Path, tree: Path) -> bool:
    """Whether `path` is `tree` or inside it, compared case-folded: macOS will
    happily write `Data-Room/` into `data-room/`."""
    inner = [part.casefold() for part in path.resolve().parts]
    outer = [part.casefold() for part in tree.resolve().parts]
    return inner[: len(outer)] == outer


def key_side(raw: str, section_dirs: Sequence[str], markers: Set[str]) -> bool:
    """Whether a cited path came from the answer-key side of a room: a folder
    before its section folder is the key root, the flagged tree, or any folder
    named for a flagged copy (`data-room-subset-flagged`). Normalising would
    otherwise throw that evidence away."""
    sections = set(section_dirs)
    for part in raw.strip().replace("\\", "/").split("/"):
        if part in sections:
            return False
        folded = part.casefold()
        if folded in markers or "flagged" in folded:
            return True
    return False


def cut_documents(cut_root: Path) -> Set[str]:
    """The documents in the uploaded cut, as room-relative `.md` paths.

    Dotfiles and dotted folders are not documents: the subset and flagged
    trees carry an ownership marker at their root, and macOS leaves
    `.DS_Store` wherever a folder was opened."""
    documents: Set[str] = set()
    for path in cut_root.rglob("*"):
        if not path.is_file():
            continue
        rel = path.relative_to(cut_root)
        if any(part.startswith(".") for part in rel.parts):
            continue
        documents.add(as_markdown(rel.as_posix()))
    return documents


# The files-read list, which must never reach the parser: glued to the last
# issue it would cite every document in the room and pre-match every finding.
# A model will not always call it "Files read", so the phrase is matched
# loosely: files or documents, then read, reviewed or opened.
_HEADING = re.compile(r"^(#{1,6})[ \t]*(.*?)[ \t]*#*[ \t]*$")
_FILES_PHRASE = re.compile(r"\b(?:files|documents)[ \t]+(?:read|reviewed|opened)\b", re.IGNORECASE)
# A level 2-4 heading, or a line on its own, that is only the label: the
# phrase, an optional bracketed note, an optional colon, optional emphasis.
# Only the label, so an issue titled "Documents reviewed by the board were
# unsigned" stays an issue.
_FILES_LABEL = re.compile(
    r"^[*_]*(?:files|documents)[ \t]+(?:read|reviewed|opened)[*_]*[ \t]*(?:\([^)\n]*\))?[ \t]*:?[*_]*$",
    re.IGNORECASE,
)
_FIRST_ISSUE = re.compile(r"^#{2,4}(?!#)", re.MULTILINE)
# `Model: x`, `**Model:** x`, `- **Skill**: x` — the hand-back asks for the
# plain form, and a model reaching for emphasis should not lose the line.
_FIELD = re.compile(
    r"^[ \t]*(?:[-*][ \t]+)?[*_]*(model|skill)[*_]*[ \t]*:[*_]*[ \t]*(.*?)[ \t]*$",
    re.IGNORECASE | re.MULTILINE,
)
# The next field's label inside a field's value: Word joins consecutive lines
# into one paragraph, so "Model: not known Skill: Built from …" arrives whole.
_INLINE_FIELD = re.compile(r"[ \t]+[*_]*(model|skill)[*_]*[ \t]*:[*_]*[ \t]*", re.IGNORECASE)
_HEADING_STYLE = re.compile(r"^heading\s+([1-9])$", re.IGNORECASE)

# Model lines that mean the agent did not know. The skill asks for "not known":
# the probe found that an example identifier in a skill is handed straight back.
_NOT_STATED = {"", "not known", "unknown", "not stated", "n/a", "none"}


@dataclass(frozen=True)
class SplitReport:
    """A report cut into the parts the importer treats differently.

    `issues` runs from the first `##`–`####` heading and is all the parser
    sees. `files_read` is None when the report has no files-read section,
    which the summary says out loud rather than treating as "read nothing".
    `set_aside` names any level-1 section after the issues that is not a
    files-read list: the hand-back allows none, and its paths must not be
    credited to the last issue.
    """

    preamble: str
    issues: str
    files_read: Optional[str]
    model: str
    skill: str
    set_aside: List[str]


def split_report(text: str) -> SplitReport:
    """Cut the files-read list, and any other trailing section, away from the
    issues.

    - A level-1 heading after the first issue ends the issues: the hand-back
      puts only `# Files read` there. If it names files read (in any wording
      `_FILES_PHRASE` knows) it is a files-read list; otherwise it is set
      aside. Every such section runs to the next level-1 heading, so two
      lists are both read.
    - A level 2-4 heading, or a line on its own, that is only the files-read
      label starts a list too. The heading's list runs to the next heading of
      its level or higher; the line's to the next heading of any level.
    """
    kept: List[str] = []
    files: List[str] = []
    set_aside: List[str] = []
    files_seen = seen_issue = False
    mode, ends_at = "keep", 0  # a files or aside section ends at a heading of level <= ends_at
    for line in text.split("\n"):
        heading = _HEADING.match(line)
        level = len(heading.group(1)) if heading else 0
        title = heading.group(2) if heading else ""
        if mode != "keep" and heading and level <= ends_at:
            mode = "keep"
        if mode == "files":
            files.append(line)
            continue
        if mode == "aside":
            continue
        if heading and level == 1 and _FILES_PHRASE.search(title) and (seen_issue or _FILES_LABEL.match(title)):
            mode, ends_at, files_seen = "files", 1, True
        elif heading and level == 1 and seen_issue:
            mode, ends_at = "aside", 1
            set_aside.append(line.strip())
        elif heading and 2 <= level <= 4 and _FILES_LABEL.match(title):
            mode, ends_at, files_seen = "files", level, True
        elif not heading and seen_issue and _FILES_LABEL.match(line.strip()):
            mode, ends_at, files_seen = "files", 6, True
        else:
            if heading and 2 <= level <= 4:
                seen_issue = True
            kept.append(line)
    remaining = "\n".join(kept)
    first = _FIRST_ISSUE.search(remaining)
    preamble = remaining[: first.start()] if first else remaining
    issues = remaining[first.start():] if first else ""
    fields = {}
    for match in _FIELD.finditer(preamble):
        value = match.group(2)
        joined = _INLINE_FIELD.search(value)
        if joined:
            fields.setdefault(joined.group(1).lower(), value[joined.end():].strip().strip("*_").strip())
            value = value[: joined.start()]
        fields[match.group(1).lower()] = value.strip().strip("*_").strip()
    return SplitReport(
        preamble=preamble,
        issues=issues,
        files_read="\n".join(files) if files_seen else None,
        model=fields.get("model", ""),
        skill=fields.get("skill", ""),
        set_aside=set_aside,
    )


def _docx_to_markdown(path: Path) -> Tuple[str, List[str]]:
    """A Word report as Markdown: `Title` and `Heading 1`–`9` become `#` runs
    (5 and deeper stay too deep for the parser, as in Markdown), list
    paragraphs become `- ` lines, and everything else is its text. Tables are
    not read — the hand-back asks for none — and are counted in a note so
    their loss is said, not silent."""
    try:
        import docx
    except ImportError:
        raise LegoraImportError(
            "reading a Word report needs python-docx — install the docx extra: "
            "pip install -e '.[docx]'"
        ) from None
    try:
        document = docx.Document(str(path))
    except Exception as exc:  # python-docx raises its own types for a bad package
        raise LegoraImportError(f"{path} could not be read as a Word document: {exc}") from None
    lines: List[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        # A link's target can be the only place a path is written: the
        # paragraph's text holds the link's words, not where it points.
        targets = [
            link.address
            for link in getattr(paragraph, "hyperlinks", [])
            if link.address and link.address not in text
        ]
        if targets:
            text = f"{text} ({', '.join(targets)})".strip()
        style = (paragraph.style.name if paragraph.style is not None else "") or ""
        heading = _HEADING_STYLE.match(style.strip())
        if not text:
            lines.append("")
        elif style.strip().lower() == "title":
            lines.append(f"# {text}")
        elif heading:
            lines.append("#" * int(heading.group(1)) + f" {text}")
        elif "list" in style.lower():
            lines.append(f"- {text}")
        else:
            lines.append(text)
    notes = []
    if document.tables:
        notes.append(
            f"the Word report has {len(document.tables)} table(s); their contents were "
            "not read — the hand-back asks for none"
        )
    return "\n".join(lines) + "\n", notes


def read_report(path: Path) -> Tuple[str, List[str]]:
    """The report as Markdown with LF line endings, and any reader notes.

    `utf-8-sig` drops a byte-order mark, which would otherwise sit in front of
    a first `##` and hide it; UTF-16, which Word's "save as text" writes, is
    read by its mark; anything else not UTF-8 is refused rather than guessed
    at. Line endings are unified because the heading patterns anchor on `$`,
    which does not match before a `\\r`."""
    suffix = path.suffix.lower()
    if suffix == ".md":
        data, notes = path.read_bytes(), []
        if data.startswith((b"\xff\xfe", b"\xfe\xff")):
            text = data.decode("utf-16")  # Word's "save as text"; the mark says which way round
        else:
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError as exc:
                raise LegoraImportError(
                    f"{path.name} is not UTF-8 text (byte {exc.start} is {data[exc.start]:#04x}). "
                    "Save the report as UTF-8 Markdown, or as Word, and import it again."
                ) from None
    elif suffix == ".docx":
        text, notes = _docx_to_markdown(path)
    else:
        raise LegoraImportError(
            f"{path.name}: a Legora review is read as Markdown (.md) or Word (.docx), "
            f"not {path.suffix or 'a file with no extension'}"
        )
    return text.replace("\r\n", "\n").replace("\r", "\n"), notes


def tool_name(model: str, skill: str, override: Optional[str] = None) -> str:
    """The `tool` field: `--tool` when given, else `legora/<model>` from the
    report's Model line, with the skill build in brackets when the report
    names it — the only record in the output of which build ran."""
    if override:
        return override
    stated = model.strip().rstrip(".").strip()
    name = "legora/" + (stated if stated.lower() not in _NOT_STATED else "model-not-stated")
    return f"{name} ({skill})" if skill else name


@dataclass(frozen=True)
class ImportResult:
    """The tool output an import produced, and the facts its summary reports."""

    output: dict
    report: Path
    room: Path
    cut: str
    cut_total: int
    by_severity: Dict[str, int]
    documents_cited: int
    files_read: Optional[int]
    not_read: int
    unexplained: List[str]
    cited_not_read: List[str]
    ignored: List[str]
    dropped: List[str]
    uncited: List[str]
    provenance: str
    notes: List[str]

    @property
    def findings(self) -> int:
        return len(self.output["findings"])


def _named(paths: Sequence[str]) -> List[str]:
    lines = [f"  {path}" for path in paths[:LIST_LIMIT]]
    if len(paths) > LIST_LIMIT:
        lines.append(f"  … and {len(paths) - LIST_LIMIT} more")
    return lines


def _unknown_message(unknown: Sequence[str], cut: str, room: Path) -> str:
    return "\n".join(
        [f"the report names {len(unknown)} path(s) that {cut}/ in {room} does not hold:"]
        + _named(unknown)
        + [
            "Nothing was written. The usual causes, most likely first: --cut names a "
            "different folder from the one uploaded to Legora; the Legora project holds "
            "a stale copy of the room; or the agent invented or mistyped a path. If it is "
            "the last, --drop-unknown imports the rest and leaves the run UNVERIFIED."
        ]
    )


def _provenance(room: Path, conf: RoomConf, unknown: Sequence[str], preamble: str) -> Tuple[str, str]:
    """The room_hash to stamp, and one sentence saying why or why not.

    The stamp vouches that this report was made on this room, and Legora never
    saw the manifest, so two things must vouch together. Every path the report
    names is in the room — necessary, but not enough: rooms are built from
    shared slot names, so every path of a smaller room exists in a larger one,
    and a path check alone verified a Quern review scored against Tarnwold's
    key. And the report's title names the room: the skill puts the Legora
    project's name there, and the README names projects after their room."""
    manifest = PurePosixPath(conf.get_relative_path("KEY_ROOT")) / MANIFEST_NAME
    if unknown:
        return "", (
            f"not stamped — {len(unknown)} path(s) the report names are not in the cut "
            "and were dropped, so it cannot be vouched for as a run on this room. The "
            "scorecard will say UNVERIFIED."
        )
    name = re.sub(r"^project\s+", "", conf.get("ROOM_CODENAME").strip(), flags=re.IGNORECASE)
    if not re.search(rf"\b{re.escape(name)}\b", preamble, re.IGNORECASE):
        return "", (
            f"not stamped — the report's title does not name {name}, and rooms share their "
            "folder names, so the paths alone cannot tie it to this room. Name the Legora "
            "project after the room and its cut (legora/README.md); the skill copies that "
            "name into the title. The scorecard will say UNVERIFIED."
        )
    content_hash = read_content_hash(room / manifest)
    if not content_hash:
        return "", (
            f"not stamped — no content_hash in {manifest} (it is written by /vdr-package). "
            "The scorecard will say UNVERIFIED."
        )
    return content_hash, (
        f"room_hash stamped from {manifest}: every path the report names is in this room, "
        f"and its title names {name}."
    )


def import_review(
    report: Path,
    room: Path,
    cut: Optional[str] = None,
    tool: Optional[str] = None,
    drop_unknown: bool = False,
    allow_wide: bool = False,
) -> ImportResult:
    """Read a Legora review and build the tool output for `score`.

    `cut` is the folder uploaded to Legora, relative to the room; it defaults
    to BLIND_TREE and cannot be inferred, because a subset run and a
    full-room run cite identical paths. Every path the report names — in an
    issue or in the files-read list — must be a document in the cut, or the
    import is refused (LegoraImportError), unless `drop_unknown`. A path from
    the answer-key side of the room is refused always: that run was not blind.
    An issue citing more than MAX_CITATIONS documents is refused unless
    `allow_wide`. Documents the files-read list leaves out are reported, never
    refused: the output is a true record of what was found either way."""
    conf = load_room_conf(room / "room.conf")
    section_dirs = conf.get_list("SECTION_DIRS")
    cut_name = (cut or conf.get("BLIND_TREE")).strip("/")
    cut_root = room / cut_name
    if not cut_root.is_dir():
        raise LegoraImportError(
            f"there is no folder {cut_name!r} in {room}. --cut names the folder that was "
            "uploaded to Legora, relative to the room: data-room, subset, data-room-pdf "
            "or data-room-docx."
        )

    text, notes = read_report(report)
    split = split_report(text)
    issues = normalise_citations(split.issues, section_dirs)
    try:
        parsed = parse_markdown_report(issues)
    except ToolOutputError:
        if split.files_read is not None and _FIRST_ISSUE.search(split.files_read):
            raise LegoraImportError(
                f"{report.name}'s files-read section comes before its issues and runs over "
                "them, so no issues are left to score. '# Files read' must come last, as the "
                "hand-back says: move it to the end of the report and import again."
            ) from None
        raise LegoraImportError(
            f"{report.name} has no issues under '##' headings, so there is nothing to "
            "score. The hand-back puts each issue under its own '##' heading; a report "
            "in another shape has to be fixed at the skill, not guessed at here."
        ) from None

    markers = {part.casefold() for part in PurePosixPath(conf.get_relative_path("KEY_ROOT")).parts}
    markers.add(PurePosixPath(conf.get_relative_path("FLAGGED_TREE")).name.casefold())
    written = citation_pairs(split.issues, section_dirs)
    if split.files_read is not None:
        written += citation_pairs(split.files_read, section_dirs)
    answers = sorted({raw for raw, _ in written if key_side(raw, section_dirs, markers)})
    if answers:
        raise LegoraImportError(
            "\n".join(
                [f"the report cites {len(answers)} path(s) from the answer-key side of the room:"]
                + _named(answers)
                + [
                    "The Legora project held the answers, so this was not a blind review and "
                    "cannot be scored. Upload only a blind cut — data-room/, subset/, "
                    "data-room-pdf/ or data-room-docx/ — to a fresh project (legora/README.md)."
                ]
            )
        )

    # The cut must be the folder Legora saw, or every coverage figure is
    # counted against the wrong one. Two tells: the tree a path names before
    # its section folder, and a rendered file type the cut does not hold.
    trees = room_trees(room, section_dirs)
    named = sorted({t for t in (tree_named(raw, section_dirs, trees) for raw, _ in written) if t})
    if named and cut_name.casefold() not in {tree.casefold() for tree in named}:
        raise LegoraImportError(
            f"the report's paths name {', '.join(t + '/' for t in named)} but --cut is "
            f"{cut_name}/, so what was and was not read would be counted against the wrong "
            f"folder. Pass --cut {named[0]}."
        )
    kinds = {PurePosixPath(raw.strip()).suffix.lower() for raw, _ in written} & set(RENDERED_SUFFIXES)
    held = {p.suffix.lower() for p in cut_root.rglob("*") if p.is_file() and not p.name.startswith(".")}
    unheld = sorted(kinds - held)
    if unheld:
        holders = sorted(t for t in trees if any(next((room / t).rglob(f"*{k}"), None) for k in unheld))
        hint = f"Pass --cut {holders[0]}." if holders else "Pass --cut with the folder that was uploaded."
        raise LegoraImportError(
            f"the report cites {' and '.join(unheld)} files, but {cut_name}/ holds none: "
            f"Legora read a rendered cut of the room. {hint}"
        )

    documents = cut_documents(cut_root)
    sections = set(section_dirs)
    cited = citations_in(issues, section_dirs)
    read_list: List[str] = []
    ignored: Set[str] = set()
    if split.files_read is not None:
        for path in citations_in(split.files_read, section_dirs):
            if path.split("/", 1)[0] in sections:
                read_list.append(path)
            else:
                ignored.add(path)
    unknown = sorted({path for path in cited + read_list if path not in documents})
    if unknown and not drop_unknown:
        raise LegoraImportError(_unknown_message(unknown, cut_name, room))

    findings = []
    for finding in parsed.findings:
        kept = [d for d in dict.fromkeys(finding.documents) if d in documents]
        findings.append(
            {
                "title": finding.title,
                "severity": finding.severity,
                "documents": kept,
                "summary": finding.summary,
            }
        )
    wide = [(f["title"], len(f["documents"])) for f in findings if len(f["documents"]) > MAX_CITATIONS]
    if wide and not allow_wide:
        raise LegoraImportError(
            "\n".join(
                [
                    f"{len(wide)} issue(s) cite more documents than any evidence chain in a "
                    f"synthetic room (the most so far is five; the limit here is {MAX_CITATIONS}):"
                ]
                + [f"  '{title}' cites {count} documents" for title, count in wide[:LIST_LIMIT]]
                + [
                    "Nothing was written. The usual cause is a list of paths — a files-read "
                    "list under a heading the importer does not know — read as part of an "
                    "issue, which would pre-match every finding it touches. Check the "
                    "report's headings against the hand-back. If an issue really rests on "
                    "that many documents, --allow-wide-issues imports it."
                ]
            )
        )
    cited_known = {d for finding in findings for d in finding["documents"]}

    if split.files_read is None:
        files_read, not_read, unexplained, cited_not_read = None, 0, [], []
    else:
        read = {path for path in read_list if path in documents}
        missed = sorted(documents - read)
        reached = {str(PurePosixPath(path).parent) for path in read}
        unexplained = [path for path in missed if str(PurePosixPath(path).parent) in reached]
        cited_not_read = sorted(cited_known - read)
        files_read, not_read = len(read), len(missed)

    room_hash, provenance = _provenance(room, conf, unknown, split.preamble)
    notes = notes + [
        f"set aside the section '{heading}' after the issues: the hand-back allows only "
        "'# Files read' there, and its paths were credited to no issue"
        for heading in split.set_aside
    ]
    output = {
        "tool": tool_name(split.model, split.skill, tool),
        "room_hash": room_hash,
        "findings": findings,
    }
    return ImportResult(
        output=output,
        report=report,
        room=room,
        cut=cut_name,
        cut_total=len(documents),
        by_severity={s: sum(1 for f in findings if f["severity"] == s) for s in SEVERITIES},
        documents_cited=len(cited_known),
        files_read=files_read,
        not_read=not_read,
        unexplained=unexplained,
        cited_not_read=cited_not_read,
        ignored=sorted(ignored),
        dropped=unknown,
        uncited=[f["title"] for f in findings if not f["documents"]],
        provenance=provenance,
        notes=notes,
    )


def write_output(result: ImportResult, out: Path) -> None:
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(result.output, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def render_summary(result: ImportResult, out: Path) -> str:
    counts = ", ".join(f"{s} {result.by_severity.get(s, 0)}" for s in SEVERITIES)
    lines = [
        f"Imported {result.findings} issue(s) from {result.report.name} ({counts}), "
        f"citing {result.documents_cited} document(s).",
        f"Cut: {result.cut}/, {result.cut_total} documents.",
    ]
    if result.files_read is None:
        lines.append(
            "Files read: the report has no '# Files read' section, so nothing can be said "
            "about what Legora did not read."
        )
    else:
        lines.append(
            f"Files read: {result.files_read} of {result.cut_total} named; {result.not_read} not named."
        )
        if result.unexplained:
            lines.append(
                f"{len(result.unexplained)} of those sit in a folder the list otherwise "
                "reaches, so a part-run does not explain them:"
            )
            lines += _named(result.unexplained)
        if result.cited_not_read:
            lines.append(f"Cited but not in the files-read list: {len(result.cited_not_read)}")
            lines += _named(result.cited_not_read)
        if result.ignored:
            lines.append(
                f"Files-read entries outside the room's section folders, ignored: {len(result.ignored)}"
            )
            lines += _named(result.ignored)
    if result.dropped:
        lines.append(f"Dropped under --drop-unknown, not in {result.cut}/: {len(result.dropped)}")
        lines += _named(result.dropped)
    if result.uncited:
        lines.append(
            f"{len(result.uncited)} issue(s) cite no document in {result.cut}/, so the pre-match "
            "cannot place them; /vdr-score's adjudication will:"
        )
        lines += _named(result.uncited)
    lines.append(f"Provenance: {result.provenance}")
    lines.append(f"Tool: {result.output['tool']}")
    lines += [f"Note: {note}" for note in result.notes]
    lines += [
        f"Written to {out}. Score it with:",
        f"  python3 -m synthvdr score {out} --room {result.room}",
    ]
    return "\n".join(lines)


def check_out(out: Path, room: Path, cut: Optional[str] = None) -> None:
    """Refuse an `--out` that is a folder, or that sits inside any tree of the
    room: the uploaded cut, the blind and flagged trees, the key, and every
    other cut.

    A tool output written into the blind tree becomes a document in it and
    changes the tree's content hash — the room would no longer be the room
    its manifest certifies. Into another cut it is uploaded as a document next
    time; into the key it would sit among the answers. `eval-runs/` and the
    like stay open."""
    conf = load_room_conf(room / "room.conf")
    if out.is_dir():
        raise LegoraImportError(
            f"--out {out} is a folder; name the file to write, for example {out / 'legora.json'}"
        )
    protected = {cut or conf.get("BLIND_TREE")}
    protected |= {conf.get_relative_path(key) for key in ("BLIND_TREE", "FLAGGED_TREE", "KEY_ROOT")}
    protected |= room_trees(room, conf.get_list("SECTION_DIRS"))
    for rel in sorted(protected):
        if _within(out, room / rel):
            raise LegoraImportError(
                f"--out {out} is inside {rel}/, a tree of the room. A tool output written into "
                "a room tree changes what that tree holds, and the blind tree's content hash "
                "with it. Write it beside the trees instead, for example in eval-runs/."
            )
